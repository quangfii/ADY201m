from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json
root=Path(__file__).resolve().parents[1]
data=json.loads((root/'public/analytics.json').read_text(encoding='utf-8'))['summary']
doc=Document(); doc.core_properties.author=''; doc.core_properties.last_modified_by=''; section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7);section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(11 if name=='Normal' else 25 if name=='Title' else 15 if name=='Heading 1' else 12)
 s.paragraph_format.space_after=Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
footer=section.footer.paragraphs[0];footer.text='ADY201m  •  Tài liệu dự án  •  '
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,1)
def code(text):
 r=doc.add_paragraph().add_run(text);r.font.name='Consolas';r.font.size=Pt(10)
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF5' if ri==0 else 'FFFFFF');pr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
    for r in para.runs:r.font.size=Pt(10);r.bold=ri==0
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 # No empty paragraph after tables
doc.add_paragraph('Dự án ADY201m so sánh giá sản phẩm',style='Title')
p('Tài liệu mô tả hệ thống và hướng dẫn vận hành')
p('Cập nhật ngày 11 tháng 10 năm 2026')
h('Mục tiêu và phạm vi')
p('ADY201m thu thập giá từ WinMart và Bách Hóa Xanh theo từng giờ, lưu lịch sử vào SQL Server trên máy người dùng và tạo báo cáo so sánh. Người dùng xem giá theo sản phẩm, nhóm hàng và thời gian, đồng thời tham khảo nơi mua rẻ hơn đối với sản phẩm ghép được.')
p('Tài liệu ghi nhận chức năng đã triển khai và các điều kiện vận hành thực tế. Phiên bản hiện tại đã có phân tích giờ nên mua theo lịch sử, kèm mức độ đầy đủ của dữ liệu.')
h('Chức năng hiện có')
table(['Chức năng','Kết quả'],[
('Thu thập định kỳ','Windows Task Scheduler chạy chu kỳ mỗi giờ'),('Lưu trữ','SQL Server 2022 với lịch sử các lần thu thập'),('So sánh sản phẩm','Giá hai siêu thị và mức chênh lệch'),('So sánh nhóm hàng','Bảng và biểu đồ giá trung bình danh mục'),('Biến động giá','Lịch sử theo thời gian và thống kê min max trung bình'),('Gợi ý mua hàng','Nơi mua rẻ hơn và khung giờ giá tốt theo lịch sử')],[2.1,4.8])
h('Dữ liệu đã xác minh')
p(f"Báo cáo tại thời điểm lập tài liệu có {data['total_observations']:,} quan sát, trong đó {data['valid_observations']:,} quan sát hợp lệ. Lần kiểm tra cào thực tế đã bổ sung 1.562 sản phẩm. Số sản phẩm ghép được có thể thay đổi khi thời điểm hoặc quy cách giá khác nhau.")
doc.add_page_break()
h('Kiến trúc và luồng xử lý')
p('Windows Task Scheduler gọi src/pipeline.js. Chu kỳ phát hiện danh mục, chạy bộ cào Playwright, chuẩn hóa quan sát, lưu vào SQL Server, phân tích lịch sử rồi cập nhật dữ liệu cho giao diện web. Nếu sitemap lỗi, hệ thống dùng danh mục từ lịch sử SQL, bộ nhớ đệm hoặc cấu hình dự phòng.')
code('Lịch mỗi giờ → Cào giá → SQL Server → Phân tích → Bảng và biểu đồ')
p('src/pipeline.js dùng file khóa để tránh hai chu kỳ chạy đồng thời. Chu kỳ không lấy được sản phẩm sẽ báo lỗi. Nếu một số danh mục lỗi, các danh mục thành công vẫn có thể được lưu; cần đọc log để đánh giá độ đầy đủ.')
h('Công nghệ sử dụng')
table(['Thành phần','Công nghệ'],[('Mã nguồn','Node.js từ 24.15 và JavaScript ES modules'),('Cào dữ liệu','Playwright và Chromium'),('Cơ sở dữ liệu','Microsoft SQL Server 2022'),('Kết nối SQL','sqlcmd ODBC và Windows Authentication'),('Giao diện','HTML CSS JavaScript và biểu đồ SVG'),('Lịch chạy','Windows Task Scheduler')],[2.1,4.8])
h('Cấu trúc dữ liệu SQL')
p('Database ADY201m chứa dbo.PriceObservations và dbo.ImportMetadata. Mỗi quan sát đại diện cho một sản phẩm tại một siêu thị vào một thời điểm. Thời gian được lưu theo UTC; cộng 7 giờ để xem giờ Việt Nam.')
table(['Trường','Ý nghĩa'],[('id','Khóa SHA256 theo thời điểm siêu thị và định danh sản phẩm'),('collected_at','Thời điểm thu thập UTC'),('supermarket và product_identity','Siêu thị và định danh sản phẩm'),('name size category','Tên quy cách và nhóm hàng'),('actual_price','Giá thực tế dương bằng VND'),('product_url và match_key','Nguồn sản phẩm và khóa ghép'),('payload','Quan sát gốc dạng JSON giữ các trường bổ sung')],[2.1,4.8])
p('Ràng buộc khóa chính chống nhập trùng. Giá không hợp lệ và thời điểm không đọc được bị loại trước khi lưu. dbo.ImportMetadata ghi nhận việc nhập CSV cũ để không nhập lặp.')
doc.add_page_break()
h('Cài đặt và chạy trên máy Windows')
p('Cài Node.js tương thích, SQL Server và sqlcmd ODBC. Đăng nhập Windows bằng tài khoản có quyền truy cập database. Clone repository rồi mở terminal trong thư mục dự án.')
code('npm ci\nnpx playwright install chromium')
p('Kiểm tra config/sqlserver.json. Cấu hình hiện tại dùng server localhost và database ADY201m. Đường dẫn sqlcmd phải phù hợp máy chạy; không cần ghi mật khẩu vì dùng Windows Authentication.')
code('npm run pipeline\nnpm start')
p('pipeline thực hiện một chu kỳ đầy đủ. npm start phục vụ giao diện mặc định tại http://localhost:4173/analytics.html. Máy hiện tại đang dùng cổng 4174; có thể đặt PORT để đổi cổng.')
code("$env:PORT='4174'\nnpm start")
h('Tự động thu thập mỗi giờ')
p('Máy hiện tại đã có tác vụ ADY201m-Hourly-Prices. Tác vụ chạy Node trực tiếp với src/pipeline.js trong thư mục dự án, mỗi giờ, theo tài khoản Windows đăng nhập. Cần máy bật, có mạng và SQL Server hoạt động. Tác vụ này nằm trong Windows và không tự đi theo repository khi clone sang máy khác.')
p('Trên máy khác, tạo tác vụ tương tự với đường dẫn Node và thư mục dự án của máy đó; đặt Multiple instances thành Ignore new instance. Hoặc dùng lệnh sau để chạy một tiến trình theo dõi liên tục, thu thập ngay rồi lặp mỗi 60 phút.')
code('npm run scrape:watch')
p('Chỉ dùng một cơ chế lịch trên cùng một máy. Xem logs/scrape.log để kiểm tra chu kỳ; scripts/hourly.ps1 là lựa chọn chạy thủ công và cần sửa đường dẫn nếu chuyển máy.')
h('Xem dữ liệu bằng SSMS')
p('Kết nối localhost bằng Windows Authentication. Mở Databases → ADY201m → Tables → dbo.PriceObservations và chọn Select Top 1000 Rows. scripts/view-prices.sql có truy vấn xem giờ Việt Nam và tổng số quan sát.')
code('USE ADY201m;\nSELECT TOP (100) name, actual_price, collected_at\nFROM dbo.PriceObservations ORDER BY collected_at DESC;')
doc.add_page_break()
h('Phương pháp phân tích và giới hạn')
p('So sánh lấy quan sát mới nhất theo định danh sản phẩm tại mỗi siêu thị. Tên được chuẩn hóa và quy cách khối lượng hoặc dung tích được quy về đơn vị g hoặc ml. Thiếu quy cách thì giữ riêng; giá hai siêu thị chỉ được so sánh khi thời điểm cách nhau không quá 90 phút.')
p('Gợi ý ưu tiên nơi có giá thấp hơn trong các cặp ghép được. Ghép tên là phương pháp thực dụng, vẫn cần kiểm tra nhãn hiệu, số lượng trong gói, biến thể và quy cách. Giá còn phụ thuộc cửa hàng, khu vực và khuyến mãi; chưa tính phí giao hàng.')
p('Giá trung bình nhóm hàng có thể dựa trên các danh mục khác nhau giữa hai siêu thị. Chỉ số này mô tả nhóm hàng đang có dữ liệu, không khẳng định một siêu thị rẻ hơn cho cùng một giỏ hàng. Dữ liệu thiếu hoặc cũ cần được xem cùng thời điểm thu thập.')
h('Kiểm thử và kết quả kiểm tra')
p('npm test kiểm tra chuẩn hóa quan sát, chống trùng theo từng thời điểm, loại giá hoặc thời gian không hợp lệ, ghép 1 lít với 1.000 ml, giữ riêng quy cách chưa rõ và loại cặp có thời điểm lệch quá giới hạn. Các kiểm thử hiện đều đạt, gồm phân tích giờ Việt Nam, dữ liệu ít, ngày không chung và quy cách khác nhau.')
p('Việc chuyển dữ liệu cũ sang SQL Server đã đối chiếu đủ 6.738 bản ghi. Nhập lại cùng dữ liệu không tạo bản ghi mới. Báo cáo và danh mục đã chạy thành công từ SQL Server; kiểm tra thực tế giữ nguyên tên tiếng Việt.')
h('Phân tích khung giờ nên mua')
p('Hệ thống phân tích 30 ngày gần nhất, riêng từng sản phẩm, quy cách và siêu thị theo giờ Việt Nam. Mỗi ngày và giờ lấy giá trung vị, rồi xếp hạng giá tương đối trên các ngày chung. Khuyến nghị cần ít nhất hai giờ và ba ngày chung. Ít hơn thì chỉ tham khảo; giá gần bằng nhau thì không chọn giờ ưu tiên. Giờ chưa thu thập để trống, dữ liệu quá 48 giờ được đánh dấu cũ. Xem mục Giờ nào nên mua và hai báo cáo CSV buying_hour_summary, hourly_product_prices. Kết quả không bảo đảm giá tương lai.')
h('Quản lý mã nguồn và bàn giao')
p('Repository GitHub lưu mã nguồn, cấu hình mẫu, kiểm thử và tài liệu Word. Cơ sở dữ liệu SQL Server, log, dữ liệu thô và các thư mục triển khai tạm vẫn ở máy chạy. GitHub Projects dùng theo dõi công việc; mã nguồn và tài liệu được lưu trong repository và liên kết vào Projects.')
p('Sau khi clone sang máy mới, cần cấu hình SQL Server, cài trình duyệt Playwright và tạo lại lịch Windows. Có thể sao lưu database qua SSMS để chuyển lịch sử giá. SQLite cũ chỉ được giữ như bản dữ liệu trước chuyển đổi, ứng dụng hiện không lưu mới vào đó.')
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for paragraph in doc.paragraphs:
 for border in list(paragraph._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
out=root/'docs/ADY201m_Tai_lieu_du_an.docx';doc.save(out);print(out)
